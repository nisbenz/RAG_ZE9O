"""Generate the committed test fixtures. Run from the repo root:

    uv run python tests/fixtures/make_fixtures.py

The chunking golden file needs tessdata and the embedder tokenizer instead (see AGENTS.md):

    uv run python tests/fixtures/make_fixtures.py --chunk-golden

Needs the dev dependencies (pillow, openpyxl, python-docx, pypdf), the DejaVu Sans font
and Pillow built with libraqm (for correct Arabic shaping in arabic_notice.png).
Outputs are committed, so tests never regenerate them.
"""

from __future__ import annotations

import asyncio
import datetime as dt
import io
import json
import os
import sys
from pathlib import Path

import docx
import openpyxl
from PIL import Image, ImageDraw, ImageFont, features
from pypdf import PdfReader, PdfWriter

OUT = Path(__file__).resolve().parent
FONT = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
FIXED_DATE = dt.datetime(2026, 1, 1, 12, 0, 0)

# --- text PDF (hand-written so the text layer is exact) -------------------------

TEXT_PDF_PAGES = [
    (
        "Club Statutes",
        [
            "The MC Club is a non-profit association open to all students.",
            "Its purpose is to organise cultural, scientific and sports activities.",
            "Membership is granted after paying the annual fee of fifty dinars.",
            "Members elect a board of seven people every two years.",
            "The board meets on the first Monday of each month at six in the evening.",
            "Decisions are taken by simple majority of the members present.",
        ],
    ),
    (
        "Annual Fees",
        [
            "The annual fee is due before the end of March each year.",
            "Students pay a reduced fee of thirty dinars on presentation of their card.",
            "Late payments incur a penalty of five dinars per month of delay.",
            "Fees may be paid in cash at the office or by bank transfer.",
            "Receipts are issued by the treasurer within one week of payment.",
            "Les cotisations financent le materiel et les sorties du club.",
        ],
    ),
    (
        "Elections",
        [
            "Board elections take place in May during the general assembly.",
            "Candidates must register with the secretary at least one week ahead.",
            "Each member has one vote and proxy voting is limited to one proxy.",
            "The vote is secret and counted in public by two volunteers.",
            "Results are published on the notice board the following day.",
            "Disputes are settled by the honorary committee within ten days.",
        ],
    ),
    (
        "Library and Rooms",
        [
            "The library opens from nine in the morning to five in the afternoon.",
            "Books may be borrowed for two weeks and renewed once.",
            "Room B12 is reserved for chess training on Monday evenings.",
            "The main hall can be booked for events with two weeks notice.",
            "Guests must sign the visitor book at the entrance.",
            "Lost items are kept at the office for one month.",
        ],
    ),
]


def _pdf_escape(text: str) -> bytes:
    raw = text.encode("cp1252")
    return raw.replace(b"\\", b"\\\\").replace(b"(", b"\\(").replace(b")", b"\\)")


def _page_stream(heading: str, lines: list[str], number: int, total: int) -> bytes:
    parts = [b"BT /F2 18 Tf 72 740 Td (" + _pdf_escape(heading) + b") Tj ET"]
    parts.append(b"BT /F1 11 Tf 16 TL 72 700 Td")
    for line in lines:
        parts.append(b"(" + _pdf_escape(line) + b") Tj T*")
    parts.append(b"ET")
    header = "MC Club Bulletin"
    footer = f"Club MC - Page {number} / {total}"
    parts.append(b"BT /F1 9 Tf 72 770 Td (" + _pdf_escape(header) + b") Tj ET")
    parts.append(b"BT /F1 9 Tf 72 40 Td (" + _pdf_escape(footer) + b") Tj ET")
    return b"\n".join(parts)


def text_pdf(pages: list[tuple[str, list[str]]]) -> bytes:
    font_regular, font_bold = 3, 4
    objects: dict[int, bytes] = {
        1: b"<< /Type /Catalog /Pages 2 0 R >>",
        font_regular: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica "
        b"/Encoding /WinAnsiEncoding >>",
        font_bold: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold "
        b"/Encoding /WinAnsiEncoding >>",
        5: b"<< /Title (MC Club Handbook) /Producer (make_fixtures) >>",
    }
    kids = []
    next_id = 6
    for i, (heading, lines) in enumerate(pages, start=1):
        page_id, content_id = next_id, next_id + 1
        next_id += 2
        stream = _page_stream(heading, lines, i, len(pages))
        objects[content_id] = b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream"
        objects[page_id] = (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 3 0 R /F2 4 0 R >> >> /Contents %d 0 R >>" % content_id
        )
        kids.append(b"%d 0 R" % page_id)
    objects[2] = b"<< /Type /Pages /Kids [" + b" ".join(kids) + b"] /Count %d >>" % len(pages)

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = {}
    for num in sorted(objects):
        offsets[num] = out.tell()
        out.write(b"%d 0 obj\n" % num + objects[num] + b"\nendobj\n")
    xref = out.tell()
    size = max(objects) + 1
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % size)
    for num in range(1, size):
        out.write(b"%010d 00000 n \n" % offsets[num])
    out.write(b"trailer\n<< /Size %d /Root 1 0 R /Info 5 0 R >>\n" % size)
    out.write(b"startxref\n%d\n%%%%EOF\n" % xref)
    return out.getvalue()


# --- images ----------------------------------------------------------------------

NOTICE_LINES = [
    "NOTICE TO ALL MEMBERS",
    "The general assembly will be held",
    "on Monday 3 March at 18:00 in room B12.",
    "Please bring your membership card.",
]
SCAN_LINES = [
    "ANNUAL FEES REMINDER",
    "The annual fee is fifty dinars.",
    "Payment is due before the end of March.",
    "Students pay thirty dinars with their card.",
]
ARABIC_LINES = [
    "إعلان إلى جميع الأعضاء",
    "يعقد النادي اجتماعه العام",
    "يوم الاثنين على الساعة السادسة مساء",
    "في القاعة الكبرى",
]


def _render(lines: list[str], *, rtl: bool = False, size=(1240, 700), font_px=44) -> Image.Image:
    image = Image.new("L", size, 255)
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype(FONT, font_px)
    y = 70
    for line in lines:
        if rtl:
            kwargs = {"direction": "rtl", "language": "ar", "features": ["-liga"]}
            width = draw.textlength(line, font=font, **kwargs)
            draw.text((size[0] - 80 - width, y), line, font=font, fill=0, **kwargs)
        else:
            draw.text((80, y), line, font=font, fill=0)
        y += int(font_px * 1.9)
    return image


def _image_pdf(image: Image.Image) -> bytes:
    buf = io.BytesIO()
    image.convert("RGB").save(buf, "PDF", resolution=150)
    return buf.getvalue()


# --- writers -----------------------------------------------------------------------


def write_pdfs() -> None:
    text = text_pdf(TEXT_PDF_PAGES)
    (OUT / "text.pdf").write_bytes(text)
    (OUT / "corrupt.pdf").write_bytes(text[: len(text) // 2])

    scanned = _image_pdf(_render(SCAN_LINES, size=(1240, 1754)))
    (OUT / "scanned.pdf").write_bytes(scanned)

    mixed = PdfWriter()
    mixed.add_page(PdfReader(io.BytesIO(text)).pages[0])
    mixed.add_page(PdfReader(io.BytesIO(scanned)).pages[0])
    buf = io.BytesIO()
    mixed.write(buf)
    (OUT / "mixed.pdf").write_bytes(buf.getvalue())

    encrypted = PdfWriter(clone_from=PdfReader(io.BytesIO(text)))
    encrypted.encrypt(user_password="secret", owner_password="owner", algorithm="RC4-128")
    buf = io.BytesIO()
    encrypted.write(buf)
    (OUT / "encrypted.pdf").write_bytes(buf.getvalue())


def write_images() -> None:
    notice = _render(NOTICE_LINES)
    notice.save(OUT / "notice.png", optimize=True)
    notice.convert("RGB").save(OUT / "notice.jpg", quality=90)
    _render(ARABIC_LINES, rtl=True).save(OUT / "arabic_notice.png", optimize=True)


def write_docx() -> None:
    document = docx.Document()
    document.core_properties.created = FIXED_DATE
    document.core_properties.modified = FIXED_DATE
    document.core_properties.title = ""
    document.add_heading("Internal Regulations", level=1)
    document.add_paragraph("These regulations complete the statutes of the MC Club.")
    document.add_heading("Board", level=2)
    document.add_paragraph("The board is composed of seven elected members.")
    table = document.add_table(rows=3, cols=2)
    for row, (role, name) in enumerate(
        [("Role", "Name"), ("Chair", "Amina"), ("Treasurer", "Karim")]
    ):
        table.cell(row, 0).text = role
        table.cell(row, 1).text = name
    document.add_heading("Meetings", level=2)
    document.add_paragraph("The board meets on the first Monday of every month.")
    document.save(OUT / "doc.docx")


def write_xlsx() -> None:
    workbook = openpyxl.Workbook()
    workbook.properties.created = FIXED_DATE
    workbook.properties.modified = FIXED_DATE
    schedule = workbook.active
    schedule.title = "Schedule"
    for row in [
        ("Day", "Activity", "Room"),
        ("Monday", "Chess", "B12"),
        ("Wednesday", "Robotics", "Lab 2"),
        ("Friday", "Debate", "Main hall"),
    ]:
        schedule.append(row)
    fees = workbook.create_sheet("Fees")
    for row in [("Category", "Amount (DT)"), ("Student", 30), ("Regular", 50)]:
        fees.append(row)
    workbook.create_sheet("Empty")
    workbook.save(OUT / "sheet.xlsx")


def write_text_files() -> None:
    (OUT / "notes.md").write_text(
        "# Club Notes\n\nGeneral information about the club.\n\n"
        "## Opening Hours\n\nThe office is open from 9:00 to 17:00.\n\n"
        "## Contact\n\nWrite to the secretary for any question.\n",
        encoding="utf-8",
    )
    (OUT / "latin1.txt").write_bytes(
        (
            "Règlement intérieur du club.\n\n"
            "Les réunions ont lieu chaque lundi à dix-huit heures dans la salle B12. "
            "Les membres doivent présenter leur carte d'adhérent à l'entrée. "
            "La cotisation annuelle s'élève à cinquante dinars, payable avant la fin du mois "
            "de mars. Les étudiants bénéficient d'un tarif réduit de trente dinars.\n"
        ).encode("latin-1")
    )
    (OUT / "arabic_cp1256.txt").write_bytes(
        (
            "النظام الداخلي للنادي.\n\n"
            "تعقد الاجتماعات كل يوم اثنين على الساعة السادسة مساء في القاعة الكبرى. "
            "يجب على الأعضاء تقديم بطاقة العضوية عند الدخول. "
            "قيمة الاشتراك السنوي خمسون دينارا تدفع قبل نهاية شهر مارس. "
            "يستفيد الطلبة من اشتراك مخفض قيمته ثلاثون دينارا.\n"
        ).encode("cp1256")
    )


def write_chunk_golden() -> None:
    sys.path.insert(0, str(OUT.parents[1]))
    from mcclub_rag.ingest.settings import IngestSettings
    from tests.chunk_helpers import MCOLI_GOLDEN, mcoli_chunk_summary

    repo = OUT.parents[1]
    settings = IngestSettings(
        tessdata_prefix=os.environ.get("TESSDATA_PREFIX", repo / "models" / "tessdata"),
        chunk_tokenizer_path=os.environ.get(
            "INGEST_CHUNK_TOKENIZER_PATH", repo / "models" / "embedder" / "tokenizer.json"
        ),
    )
    summary = asyncio.run(mcoli_chunk_summary(settings))
    MCOLI_GOLDEN.parent.mkdir(exist_ok=True)
    rows = ",\n".join(f"  {json.dumps(row, ensure_ascii=False)}" for row in summary)
    MCOLI_GOLDEN.write_text(f"[\n{rows}\n]\n")  # one chunk per line: readable diffs
    print(f"{MCOLI_GOLDEN.name}: {len(summary)} chunks")


def main() -> None:
    if "--chunk-golden" in sys.argv[1:]:
        write_chunk_golden()
        return
    if not features.check("raqm"):
        raise SystemExit("Pillow needs libraqm to shape Arabic text correctly")
    write_pdfs()
    write_images()
    write_docx()
    write_xlsx()
    write_text_files()
    for path in sorted(OUT.iterdir()):
        if path.suffix != ".py" and path.name != "README.md":
            print(f"{path.name:22} {path.stat().st_size:>8} bytes")


if __name__ == "__main__":
    main()
